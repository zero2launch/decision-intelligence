import base64
import csv
import io
import json
import logging
from concurrent.futures import Future, ThreadPoolExecutor
from datetime import datetime, timezone


import pymongo.errors

from src.repository.upload import UploadRepository
from src.schemas.upload import UploadRequest, UploadResponse
from src.services.chunking import ChunkingService
from src.services.knowledge_graph import KnowledgeGraphService
from src.utils.exception import AppException

logger = logging.getLogger(__name__)


class UploadService:
    def __init__(self):
        self._repo = UploadRepository()
        self._chunking_svc = ChunkingService()
        self._kg_svc = KnowledgeGraphService()

    def process_upload(self, payload: UploadRequest, username: str) -> UploadResponse:
        uploaded_count = 0
        for file in payload.files:
            content_type = file.filename.rsplit(".", 1)[-1].lower()
            if content_type not in {"pdf", "csv", "xls", "xlsx"}:
                raise AppException(
                    f"Unsupported file type: {file.filename}. Accepted: pdf, csv, xls, xlsx",
                    status_code=400,
                )

            try:
                binary_data = base64.b64decode(file.content, validate=True)
            except Exception:
                raise AppException(
                    f"Invalid Base64 content in file {file.filename}",
                    status_code=400,
                )

            if content_type == "pdf":
                extracted_text = self._extract_pdf(binary_data)
            elif content_type == "csv":
                extracted_text = self._extract_csv(binary_data)
            elif content_type == "xls":
                extracted_text = self._extract_xls(binary_data)
            else:
                extracted_text = self._extract_xlsx(binary_data)


            if content_type == "pdf":
                with ThreadPoolExecutor(max_workers=2) as executor:
                    vector_future: Future = executor.submit(
                        self._chunking_svc.chunk_and_store,
                        extracted_text, file.filename, username,
                    )
                    kg_future: Future = executor.submit(
                        self._kg_svc.process_document,
                        extracted_text, file.filename,
                    )
                vector_future.result()
                try:
                    kg_future.result()
                except Exception as e:
                    logger.warning(f"KG pipeline failed for '{file.filename}', continuing: {e}")

            elif content_type in {"csv", "xls", "xlsx"}:
                self._kg_svc.process_document(extracted_text, file.filename)

            self._repo.insert_document({
                "filename": file.filename,
                "username": username,
                "size": file.size,
                "uploaded_at": datetime.now(timezone.utc),
            })
            uploaded_count += 1

        return UploadResponse(success=True, uploaded=uploaded_count)

    def _extract_pdf(self, binary_data: bytes) -> str:
        import fitz
        doc = fitz.open(stream=binary_data, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)

    def _extract_csv(self, binary_data: bytes) -> str:
        text = binary_data.decode("utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = [row for row in reader]
        return json.dumps(rows, ensure_ascii=False)

    def _extract_xls(self, binary_data: bytes) -> str:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=binary_data)
        rows = []
        for sheet in workbook.sheets():
            headers = [sheet.cell_value(0, col) for col in range(sheet.ncols)]
            for row_idx in range(1, sheet.nrows):
                rows.append({headers[col]: sheet.cell_value(row_idx, col) for col in range(sheet.ncols)})
        return json.dumps(rows, ensure_ascii=False)

    def _extract_xlsx(self, binary_data: bytes) -> str:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(binary_data), data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            for row in sheet.iter_rows(min_row=2, values_only=True):
                rows.append({headers[col]: row[col] for col in range(len(headers))})
        return json.dumps(rows, ensure_ascii=False)
